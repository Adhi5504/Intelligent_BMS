#!/usr/bin/env python3
"""Assemble AI_PBMS_8Slide.pptx from the generated and extracted assets."""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import theme

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GEN  = os.path.join(ROOT, "assets", "generated")
EXT  = os.path.join(ROOT, "assets", "extracted")
OUT  = os.path.join(ROOT, "AI_PBMS_8Slide.pptx")

def C(h): return RGBColor.from_string(h.lstrip("#"))
BG, PANEL, INK  = C(theme.BG), C(theme.PANEL), C(theme.INK)
MUTED, HAIRLINE = C(theme.MUTED), C(theme.HAIRLINE)
ACCENT, WARNING, CRITICAL = C(theme.ACCENT), C(theme.WARNING), C(theme.CRITICAL)

HEAD_FONT, BODY_FONT = theme.HEAD_FONT, theme.BODY_FONT

SW, SH = 13.333, 7.5                      # slide size, inches
M      = 0.62                             # outer margin
TITLE_Y, TITLE_H = 0.40, 0.78
RULE_Y = 1.26                             # thin accent rule under the title
BODY_Y = 1.56
BODY_B = 6.86                             # bottom of the content band
LEFT_W = 6.34                             # ~55% content column (net of margin)
VIS_X  = M + LEFT_W + 0.36
VIS_W  = SW - M - VIS_X                   # ~45% visual column


# ----------------------------------------------------------------- helpers
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, runs, size=17, color=INK, bold=False, font=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=0, line=1.18):
    """runs: a string, or a list of (string, {overrides}) for per-para styling."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = [runs] if isinstance(runs, str) else runs
    for i, item in enumerate(items):
        s, ov = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ov.get("align", align)
        p.line_spacing = ov.get("line", line)
        p.space_after = Pt(ov.get("space_after", space_after))
        r = p.add_run(); r.text = s
        f = r.font
        f.name = ov.get("font", font or BODY_FONT)
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.color.rgb = ov.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=17, dot_col=None, gap=9):
    """Body copy as genuine bulleted paragraphs.

    Uses real DrawingML bullet properties (buChar + buClr + hanging indent) so
    spacing stays even regardless of how many lines each bullet wraps to --
    a fixed-pitch layout drifts as soon as one bullet is longer than another.
    """
    dot_col = dot_col or ACCENT
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    marL, ind = Emu(int(0.30 * 914400)), Emu(int(-0.30 * 914400))
    for i, txt in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = 1.16
        para.space_after = Pt(gap)
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(marL)); pPr.set("indent", str(ind))
        _bullet_glyph(pPr, dot_col, size)
        r = para.add_run(); r.text = txt
        r.font.name = BODY_FONT; r.font.size = Pt(size); r.font.color.rgb = INK
    return tb


def _bullet_glyph(pPr, colour, size):
    """Attach <a:buClr>/<a:buSzPct>/<a:buFont>/<a:buChar> in schema order."""
    from pptx.oxml.ns import qn
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    clr = etree.SubElement(pPr, qn("a:buClr"))
    srgb = etree.SubElement(clr, qn("a:srgbClr")); srgb.set("val", str(colour))
    sz = etree.SubElement(pPr, qn("a:buSzPct")); sz.set("val", "62000")
    fnt = etree.SubElement(pPr, qn("a:buFont"))
    fnt.set("typeface", "Arial"); fnt.set("pitchFamily", "34"); fnt.set("charset", "0")
    ch = etree.SubElement(pPr, qn("a:buChar")); ch.set("char", "\u25cf")


def fit(path, bx, by, bw, bh):
    """Letterbox an image inside a box, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    sc = min(bw / iw, bh / ih)
    w, h = iw * sc, ih * sc
    return bx + (bw - w) / 2, by + (bh - h) / 2, w, h


def picture(slide, path, bx, by, bw, bh, card=False, pad=0.13):
    """Place an image; `card` wraps it in a white 1px #CBD5E1 panel so that
    photos and screenshots with non-white backgrounds do not clash."""
    if card:
        x, y, w, h = fit(path, bx + pad, by + pad, bw - 2*pad, bh - 2*pad)
        rect(slide, x - pad, y - pad, w + 2*pad, h + 2*pad,
             fill=BG, line=HAIRLINE, lw=1.0)
    else:
        x, y, w, h = fit(path, bx, by, bw, bh)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


def caption(slide, x, y, w, s):
    text(slide, x, y, w, 0.26, s, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def new_slide(prs, n, title, eyebrow=None):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bgf = sl.background.fill; bgf.solid(); bgf.fore_color.rgb = BG
    ty = TITLE_Y
    if eyebrow:
        text(sl, M, TITLE_Y - 0.06, LEFT_W + 3.0, 0.26, eyebrow,
             size=11, color=ACCENT, bold=True)
        ty = TITLE_Y + 0.24
    text(sl, M, ty, SW - 2*M, TITLE_H, title, size=33, color=INK,
         bold=True, font=HEAD_FONT, line=1.02)
    # thin accent rule under the title
    r = rect(sl, M, RULE_Y + (0.22 if eyebrow else 0), 2.15, 0.035, fill=ACCENT)
    r.line.fill.background()
    r2 = rect(sl, M + 2.15, RULE_Y + (0.22 if eyebrow else 0), SW - 2*M - 2.15,
              0.035, fill=HAIRLINE)
    r2.line.fill.background()
    # running footer + slide number
    text(sl, M, SH - 0.50, 5.0, 0.26, "AI-PBMS  ·  Team ANS_4X", size=9.5, color=MUTED)
    text(sl, SW - M - 1.2, SH - 0.50, 1.2, 0.26, str(n), size=10.5, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return sl


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s.strip()


def stat_strip(slide, x, y, w, items):
    """Row of small key-figure cards."""
    n = len(items); gap = 0.16
    cw = (w - gap * (n - 1)) / n
    for i, (big, cap, col) in enumerate(items):
        cx = x + i * (cw + gap)
        rect(slide, cx, y, cw, 0.86, fill=PANEL, line=HAIRLINE, lw=0.75)
        text(slide, cx, y + 0.10, cw, 0.36, big, size=16, color=col, bold=True,
             align=PP_ALIGN.CENTER)
        text(slide, cx, y + 0.50, cw, 0.28, cap, size=9, color=MUTED,
             align=PP_ALIGN.CENTER)


def table(slide, x, y, w, headers, rows, col_w, head_h=0.46, row_h=0.52,
          fs=11, hfs=10.5):
    """Hand-built table: rectangles + textboxes, so every colour is ours."""
    total = sum(col_w)
    col_w = [c / total * w for c in col_w]
    # header
    cx = x
    rect(slide, x, y, w, head_h, fill=PANEL, line=HAIRLINE, lw=0.75)
    for cwi, h in zip(col_w, headers):
        text(slide, cx + 0.09, y, cwi - 0.18, head_h, h, size=hfs, color=INK,
             bold=True, anchor=MSO_ANCHOR.MIDDLE,
             align=PP_ALIGN.LEFT if cx == x else PP_ALIGN.CENTER)
        cx += cwi
    # body
    ry = y + head_h
    for ri, row in enumerate(rows):
        if ri == 0:                                    # highlight the chosen model
            rect(slide, x, ry, w, row_h, fill=C("#E6F4F6"), line=ACCENT, lw=1.1)
        else:
            ln = rect(slide, x, ry + row_h, w, 0.012, fill=HAIRLINE)
            ln.line.fill.background()
        cx = x
        for ci, (cwi, cell) in enumerate(zip(col_w, row)):
            val, col, bold = cell if isinstance(cell, tuple) else (cell, INK, False)
            text(slide, cx + 0.09, ry, cwi - 0.18, row_h, val, size=fs, color=col,
                 bold=bold or ri == 0 and ci == 0, anchor=MSO_ANCHOR.MIDDLE,
                 align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)
            cx += cwi
        ry += row_h
    return ry


